from pathlib import Path
import csv
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
OUT=ROOT/'submission-documents'

def document(title, purpose):
    d=Document()
    s=d.sections[0]
    s.top_margin=s.bottom_margin=Inches(.7)
    s.left_margin=s.right_margin=Inches(.75)
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(10.5 if name=='Normal' else 22 if name=='Title' else 15 if name=='Heading 1' else 12)
        st.paragraph_format.space_after=Pt(4)
        st.paragraph_format.line_spacing=1
    d.styles['Normal'].paragraph_format.widow_control=True
    d.core_properties.author='Ravi'
    d.core_properties.title=title
    d.core_properties.subject='iPhone Sales Analytics Platform submission'
    d.add_paragraph(title,'Title')
    d.add_paragraph('Submitted by Ravi')
    d.add_paragraph(purpose)
    f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=f.add_run('Ravi  |  ');r.font.size=Pt(9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    return d

def p(d,text): return d.add_paragraph(text)
def h(d,text): return d.add_paragraph(text,'Heading 1')
def question(d,n,short,q,a):
    d.add_paragraph(f'{n} {short}','Heading 2')
    x=p(d,q);x.runs[0].bold=True
    p(d,a)
def code(d,text):
    x=p(d,text)
    x.paragraph_format.space_after=Pt(4)
    x.paragraph_format.line_spacing=1
    for r in x.runs:r.font.name='Consolas';r.font.size=Pt(9)
    return x
def table(d,headers,records,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,value in zip(t.rows[0].cells,headers):c.text=str(value)
    for row in records:
        for c,value in zip(t.add_row().cells,row):c.text=str(value)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1' if ri==0 else 'FFFFFF');pr.append(shade)
            for x in c.paragraphs:
                x.paragraph_format.space_after=Pt(0);x.paragraph_format.line_spacing=1.02
                if ci>0:x.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in x.runs:r.font.size=Pt(10);r.bold=(ri==0)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    p(d,'')
    return t

def logrows(n):
    text=(ROOT/'logs'/f'actual-evidence-{n:02d}.log').read_text(encoding='utf-8-sig')
    return [[c.strip() for c in line.strip().strip('|').split('|')] for line in text.splitlines() if line.startswith('|')]
products=[r for r in logrows(6) if r[0] in ['AirPods','iPhone 13','iPhone 14']]
stores=[r for r in logrows(7) if r[0].startswith('Apple Store')]
daily=[r for r in logrows(11) if r[0].startswith('2023-')]
customers=[r for r in logrows(11) if r[0] in ['Alice','Bob','Charlie']]
fact=[r for r in logrows(10) if r[0].isdigit()]
counts=[]
for name in ['customers','products','stores','sales']:
    with (ROOT/'data'/f'{name}.csv').open() as file:counts.append((name+'.csv',len(list(csv.DictReader(file)))))

qa=document('Ravi iPhone Sales Project Questions and Answers','This submission explains the project implementation, verifies the required outputs, and answers the design questions. The accompanying numbered PNG files are direct screenshots of actual Hive and HDFS commands.')
h(qa,'Project objective and input')
question(qa,1,'Business objective','What business problem does the project address?','The retailer needs a consistent view of sales by product, store, customer and date. The project converts raw CSV data into typed, validated Hive tables and a star schema for revenue and purchase reporting.')
question(qa,2,'Technology and flow','Which tools and architecture are used?','CSV files are ingested into HDFS. PySpark processes them through Bronze, Silver and Gold layers. Hive stores the table metadata, Parquet stores the warehouse files, and Hive SQL produces the business reports. The existing Docker course container provides the local environment.')
code(qa,'CSV -> HDFS /iphone/raw -> Bronze -> Silver -> Gold star schema -> Hive SQL')
question(qa,3,'Source data','Which input files and records were supplied?','The four files below reproduce the assignment sample. AirPods is included because it is part of the supplied product and sales data. The source does not specify a currency, so revenue values use the supplied price units.')
table(qa,['Source file','Records'],counts,[4.7,1.3])
question(qa,4,'Raw ingestion','How is the Bronze layer loaded?','The reader uses the CSV header and reads source fields as strings. It checks the expected column names and uses FAILFAST for malformed CSV input. Each source is written as a separate Parquet Hive table with an explicit HDFS path. There are four Bronze tables: bronze_customers, bronze_products, bronze_stores and bronze_sales.')
question(qa,5,'Project identification','Where is Ravi saved in the actual project?','The database and all 13 tables store project.student = Ravi and project.name = Ravi iPhone Sales Analytics. HDFS also stores /iphone/project-info.txt. The project files are in /home/takeo/ravi-iphone-project. The existing course account is takeo; the assignment database name remains iphone_analytics. Screenshots 2, 3, 4 and 12 show the persisted name.')

qa.add_page_break();h(qa,'Silver processing and Gold modeling')
question(qa,6,'Cleaning rules','How are raw records cleaned and standardized?','Silver trims values, casts identifiers and quantities to INT, parses dates as DATE, normalizes customer/store state abbreviations, removes null or invalid records and exact duplicates, and stops on conflicting duplicate identifiers. Sales quantity and identifiers must be positive; product price must be nonnegative. The supplied sample contains four valid sales records, so all four remain. Separate bad-input test cases are not part of the submitted evidence.')
question(qa,7,'Join validation','How are missing dimension references handled?','Before Gold is written, left-anti joins check every sale against products, customers and stores. An unmatched product_id, customer_id or store_id raises an error. This prevents unmatched sales from silently disappearing during the fact join.')
question(qa,8,'Silver tables','Which Silver tables are created and partitioned?','silver_customers, silver_products, silver_stores and silver_sales are Parquet tables. silver_sales is partitioned by sale_date. Its nonpartition columns are sale_id, product_id, customer_id, store_id and quantity, all INT; sale_date is DATE.')
question(qa,9,'Dimensions','What dimensions form the Gold star schema?','Four dimensions join to fact_sales using the supplied natural identifiers. dim_date contains the four observed sales dates, with year, month and day derived from those dates. It is not a complete calendar of dates without sales.')
table(qa,['Dimension','Attributes'],[
('dim_customer','customer_id, customer_name, city, state'),('dim_product','product_id, product_name, category, unit_price'),('dim_store','store_id, store_name, city, state'),('dim_date','date_key, year, month, day')],[1.65,4.75])
question(qa,10,'Fact grain and amount','What does one fact_sales row represent and how is revenue calculated?','One row represents one supplied sale record. Sales joins to products by product_id. total_amount is quantity multiplied by unit_price. The fact stores sale_id, customer_id, product_id, store_id, quantity and total_amount, and is partitioned by date_key. The first five numeric columns are INT; total_amount is BIGINT and date_key is DATE. For sale 2, 2 times 1100 equals 2200.')

qa.add_page_break();h(qa,'Mandatory execution checks')
question(qa,11,'Database and tables','Does Hive contain the required database and all layer tables?','Yes. Screenshot 1 shows iphone_analytics and 13 tables: four Bronze, four Silver, four dimensions and fact_sales. The Gold naming requirement allows fact_* and dim_* names; a separate gold_* prefix is not necessary.')
code(qa,'SHOW DATABASES;\nUSE iphone_analytics;\nSHOW TABLES;')
question(qa,12,'Storage and schema','Do the Silver and fact schemas show Parquet, date partitions and HDFS locations?','Yes. Screenshots 2 and 3 show Parquet input/output formats, sale_date or date_key as DATE partition columns, and full HDFS locations ending in /silver_sales and /fact_sales. They also show Ravi in the stored table properties.')
code(qa,'DESCRIBE FORMATTED silver_sales;\nDESCRIBE FORMATTED fact_sales;')
question(qa,13,'Medallion directories','Does HDFS contain the required layer directories?','Yes. Screenshot 4 shows 13 warehouse directories, including bronze_sales, silver_sales, fact_sales and dim_product. The HDFS project information file identifies Ravi.')
code(qa,'hadoop fs -ls /warehouse/tablespace/managed/hive/')
question(qa,14,'Partition verification','Which date partitions were actually written?','Both sales tables have four partitions: 2023-01-10, 2023-01-12, 2023-01-15 and 2023-01-20. Silver uses sale_date and the fact uses date_key. Screenshot 5 shows both listings; the samples and counts confirm that the tables contain data.')
code(qa,'SHOW PARTITIONS silver_sales;\nSHOW PARTITIONS fact_sales;')
question(qa,15,'Row counts','Do the counts prove that the valid sample sales were preserved?','Yes. Screenshot 8 shows 4 rows in bronze_sales, 4 in silver_sales and 4 in fact_sales. The COUNT queries include descriptive aliases but have the same meaning as the assignment queries. DESCRIBE statistics may show numRows = 0 because Spark has not populated Hive statistics; actual COUNT results are the data check.')
code(qa,'SELECT COUNT(*) FROM bronze_sales;\nSELECT COUNT(*) FROM silver_sales;\nSELECT COUNT(*) FROM fact_sales;')

qa.add_page_break();h(qa,'Business query answers')
question(qa,16,'Product revenue','What revenue does each product generate?','The actual Hive query joins fact_sales to dim_product and groups by product_name. Screenshot 6 contains the query, column names and the following output rows.')
code(qa,'SELECT p.product_name, SUM(f.total_amount) AS revenue\nFROM fact_sales f\nJOIN dim_product p ON f.product_id = p.product_id\nGROUP BY p.product_name;')
table(qa,['Product','Revenue'],products,[4.4,1.6])
question(qa,17,'Store revenue','What revenue does each store generate?','The actual Hive query joins the store dimension to the fact. Screenshot 7 confirms that Apple Store NY generates 3100 and Apple Store SJ generates 1300.')
code(qa,'SELECT s.store_name, SUM(f.total_amount) AS revenue\nFROM fact_sales f\nJOIN dim_store s ON f.store_id = s.store_id\nGROUP BY s.store_name;')
table(qa,['Store','Revenue'],stores,[4.4,1.6])
question(qa,18,'Revenue integrity','Do the product, store and fact totals agree?','Yes. Product totals are 400 + 1800 + 2200 = 4400. Store totals are 3100 + 1300 = 4400. The four fact amounts are 900, 2200, 900 and 400, which also total 4400. Total quantity is 1 + 2 + 1 + 2 = 6 units. These are totals from the supplied sample, not a claim about a larger retail dataset.')

qa.add_page_break();h(qa,'Time reports customers and data samples')
question(qa,19,'Daily and monthly reporting','What are the daily and monthly sales results?','Screenshot 11 shows these daily quantities and revenue. Joining dim_date to fact_sales gives year 2023, month 1 and monthly revenue 4400. The date range is January 10 to January 20, 2023, with sales on four dates.')
table(qa,['Date','Units','Revenue'],daily,[3,1.3,1.7])
question(qa,20,'Customer behavior','What does the sample show about customer purchases?','Alice has two sale records and buys three units. Bob has one sale record and buys two units; his revenue contribution is the highest at 2200. Charlie has one record and buys one unit. Screenshot 11 shows the grouped customer output.')
table(qa,['Customer','Purchases','Units','Revenue'],customers,[2.3,1.2,1.1,1.4])
question(qa,21,'Sampling proof','Do the actual Silver and fact samples match the input?','Yes. Screenshots 9 and 10 show the four sale identifiers and their matching customer, product, store and date values. The fact adds the correctly calculated amount. The commands use LIMIT 10; four rows are returned because only four records were supplied.')
code(qa,'SELECT * FROM silver_sales LIMIT 10;\nSELECT * FROM fact_sales LIMIT 10;')
table(qa,['Sale','Cust ID','Product','Store','Qty','Amount','Date'],fact,[.45,.65,.65,.55,.45,.8,1.2])

qa.add_page_break();h(qa,'Interview questions and implementation notes')
question(qa,22,'Parquet','Why Parquet?','Parquet stores data by column and supports compression. Analytical queries can read the needed columns rather than scanning every field in CSV text. The submitted schemas confirm Parquet storage. This project does not include a timed CSV-versus-Parquet benchmark.')
question(qa,23,'Medallion architecture','Why Medallion Architecture?','Bronze preserves the source, Silver creates consistent records, and Gold presents business measures and dimensions. These separate stages make it easier to inspect data, reuse clean tables and locate a problem without mixing raw ingestion with reporting logic.')
question(qa,24,'Star schema','Why Star Schema in Hive?','A central sales fact with customer, product, store and date dimensions gives reporting queries clear joins and reusable descriptive attributes. It is suitable for BI-style aggregation. Actual speed depends on data size and the query engine; the small sample does not establish a performance benchmark.')
question(qa,25,'Date partitions','Why partition fact by date?','Time-filtered reports can limit the partitions they read. This supports partition pruning and organizes the stored files by date. The submitted fact has four date_key partitions. Full-total queries still need all relevant dates.')
question(qa,26,'Reusable code','How is code reused and what is submitted?','pipeline.py contains a reusable write() function and a schemas dictionary, plus the source-specific cleaning, validation, dimension and fact logic. The package includes the CSVs, SQL, README, documentation and actual screenshots. No shared GitHub repository was created. The suggested common-utils/bronze/silver/gold directory layout is a recommendation; this sample uses a single pipeline module and a separate SQL directory.')
question(qa,27,'Reference differences','Are there any differences from the reference implementation?','Yes. Explicit path-backed Parquet tables are EXTERNAL_TABLE rather than the reference examples\' default managed tables. total_amount is BIGINT rather than INT to avoid integer multiplication overflow. date_key is declared once as a typed partition column, correcting the reference fact DDL that lists it twice and omits its partition datatype. The required table names, joins, amount formula, date partitions, HDFS storage and sample outputs match.')
question(qa,28,'Scope and limits','What remains outside this sample demonstration?','The load overwrites a complete snapshot, as the reference examples do. Production daily operation would need incremental ingestion, scheduling, lineage, rejected-record reporting and historical pricing. Cleaning logic is implemented, but the submitted execution evidence covers the supplied valid sample. A final grade is determined by the instructor; this verification confirms the functional checks, not a guaranteed score.')

qa.add_page_break();h(qa,'Submission requirement verification')
p(qa,'The mandatory screenshot requirements were checked against the supplied assignment, the actual console transcripts, the PNG files and the sample calculations.')
table(qa,['Assignment requirement','Result','Files'],[
('Database and Bronze Silver Gold tables','Matches','1.png'),
('Silver and fact schemas with Parquet HDFS and date partitions','Matches','2.png, 3.png'),
('Required HDFS warehouse directories','Matches','4.png'),
('Both date partition listings','Matches','5.png'),
('At least two SQL queries with named output columns','Matches','6.png, 7.png'),
('All three sales counts','Matches at 4 each','8.png'),
('Silver and fact LIMIT samples','Matches','9.png, 10.png'),
('Daily monthly and customer reporting','Verified','11.png'),
('Ravi saved in actual project properties','Verified','12.png')],[3.85,1.2,1.25])
p(qa,'The mandatory numerical and execution results match. The external-table choice, BIGINT amount type, snapshot scope and recommended-folder variation are explicitly explained above. Keep the original numbered PNG files with this Word document so their text can be viewed at full size. Credentials are stored separately and are not part of the submission archive.')
qa.save(OUT/'Ravi_Questions_and_Answers.docx')

report=document('Ravi iPhone Sales Analytics Project Report','This report describes my project, the checks performed on the supplied sales data and the business results produced by Hive.')
h(report,'Project overview')
p(report,'My objective was to turn the retailer\'s CSV files into tables that could answer sales questions consistently. The input contained three customers, three products, two stores and four sales records. I used the existing Docker course environment for HDFS, PySpark and Hive. The project is named Ravi iPhone Sales Analytics in the actual Hive properties and in the HDFS project information file.')
p(report,'The final sample contains four fact records, six units and total revenue of 4400. Product, store and monthly totals agree with that amount. I kept all supplied records, including the AirPods sale, because it belongs to the reference dataset.')
h(report,'How the data moves through the project')
p(report,'The CSV files are stored in HDFS at /iphone/raw. Bronze keeps the incoming fields as strings and writes one Parquet table for each file. This gives the later stages a consistent raw source to work from.')
p(report,'Silver prepares the data for joins. The code trims text, casts identifiers and quantities, parses dates and normalizes state abbreviations. It removes invalid or duplicate records and checks for conflicting keys. Before loading the fact, it checks that every sale has a matching customer, product and store. The supplied records are valid, so the count stays at four. I have not included a separate run with deliberately bad input.')
p(report,'Gold uses a sales fact joined to four dimensions. Each fact row represents a sale record. Product prices are joined to sales, and quantity times unit price gives total_amount. The dimensions hold customer, product, store and date attributes. Silver sales is partitioned by sale_date, and fact sales is partitioned by date_key. Both tables have four January 2023 partitions.')
h(report,'Why I used this design')
p(report,'The three layers separate raw ingestion, data preparation and reporting. Parquet is suitable for column-based analytical reads, while the star schema gives the SQL queries clear joins. Date partitions organize the files and can reduce the amount read by date-filtered queries. These choices are supported by the stored schemas; the sample is too small to claim a measured performance improvement.')

report.add_page_break();h(report,'Sales findings')
p(report,'The iPhone 14 sale contributes 2200, the two iPhone 13 sales contribute 1800, and AirPods contributes 400. The product totals sum to 4400. Apple Store NY contributes 3100, compared with 1300 from Apple Store SJ. These are results from the four supplied records, so they describe this sample rather than a wider business trend.')
table(report,['Product','Revenue'],products,[4.4,1.6])
table(report,['Store','Revenue'],stores,[4.4,1.6])
p(report,'All sales fall in January 2023, giving monthly revenue of 4400. January 12 has the largest daily amount, 2200, from Bob\'s two-unit iPhone 14 purchase. Alice has two purchase records and buys three units; Bob buys two units in one record; Charlie buys one unit in one record.')
table(report,['Customer','Purchases','Units','Revenue'],customers,[2.3,1.2,1.1,1.4])
h(report,'Execution checks')
p(report,'I checked the database and all 13 table names, inspected the Silver and fact schemas, listed the HDFS warehouse directories and verified both sets of date partitions. Hive returned four rows for Bronze sales, Silver sales and fact sales. The data samples show the same sale identifiers and matching attributes, with the calculated fact amounts.')
p(report,'The numbered PNG files are direct captures of real console windows running the commands. Files 1 to 10 cover the mandatory checks, file 11 contains the additional reports, and file 12 shows Ravi saved in the database and table properties. The existing Linux account is takeo; the project\'s student name is Ravi. No separate Linux account was needed.')

report.add_page_break();h(report,'Implementation decisions and practical limits')
p(report,'There are two deliberate storage and type differences from the reference examples. The tables use explicit external HDFS locations, and the fact amount uses BIGINT so larger integer calculations do not overflow INT. The fact partition column is declared once with a DATE type. The table names, sales joins, amount calculation and required date partitions match the assignment.')
p(report,'The course image runs through ARM emulation on this machine. A query-service memory limit caused an initial parallel attempt to fail. The project service was adjusted to a 2 GB heap and a 16 MB sorting buffer, and the final captured queries completed successfully. Java subprocesses use FORK, and the optional Spark UI is disabled to avoid the emulated process and shutdown issues observed during execution.')
p(report,'The current load is a complete-snapshot overwrite. That follows the sample approach, but it is not a production daily ingestion system. For a larger retailer, I would add incremental loading, an audit trail for each batch, rejected-record storage, scheduling and historical product prices. A complete calendar dimension would allow reporting on dates with no sales. Decimal currency types would be needed if the prices contain fractions.')
h(report,'Submission and verification outcome')
p(report,'The mandatory output checks match the assignment and the supplied sample calculations. The code, CSVs, SQL, numbered screenshots and documentation are included in the project submission. The question-and-answer document explains the design questions and the implementation differences. Private connection details are kept in a separate credentials file and are excluded from the submission ZIP.')
p(report,'The result is a working sample analytics platform with consistent sales totals and traceable execution output. The evidence demonstrates the four valid source records; it does not claim a larger dataset, a production scheduler, invalid-data test coverage or guaranteed marks.')
report.save(OUT/'Ravi_Project_Report.docx')
print('Created two Word submission documents')

